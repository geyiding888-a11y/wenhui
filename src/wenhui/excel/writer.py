"""输出：生成「汇总表」和「问题清单」两个 Excel。

## 汇总表的设计取舍

汇总表是你要**直接交上去**的，所以主体必须跟标准表一模一样，不能有多余的列。

但出了问题时你又需要知道"这行是哪份表来的"。折中办法：
把「来源文件」「来源行」两列**放在最后、涂成灰色**——
一眼能看出是程序加的辅助信息，交之前删掉这两列即可。

## 问题清单的设计

一条问题要能让**不看你表格的人**直接去问填报人。所以每条都带齐：
哪个文件、哪个工作表、第几行、哪一列、现在填的是什么、建议怎么改。
按"必须处理 → 请确认 → 仅供参考"排序，最要紧的排最前面。
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .validator import Issue, Record

# ---------------------------------------------------------------- 样式常量

_HEADER_FONT = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_BODY_FONT = Font(name="微软雅黑", size=10)

_AUX_FONT = Font(name="微软雅黑", size=10, color="808080")
_AUX_FILL = PatternFill("solid", fgColor="F2F2F2")

_UNCERTAIN_FILL = PatternFill("solid", fgColor="FFF2CC")   # 淡黄：AI 拿不准的列

_LEVEL_STYLE = {
    "error": (PatternFill("solid", fgColor="FFC7CE"), Font(name="微软雅黑", bold=True, color="9C0006")),
    "warning": (PatternFill("solid", fgColor="FFEB9C"), Font(name="微软雅黑", color="9C6500")),
    "info": (PatternFill("solid", fgColor="DDEBF7"), Font(name="微软雅黑", color="1F4E79")),
}

_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_MAX_COL_WIDTH = 42
_MIN_COL_WIDTH = 8

#: 汇总表末尾附带的溯源列（灰色，可删）。
#: **公开的**：问数助手建查询表时要用同一份列名——
#: 你在 Excel 里看到的列名，和它眼里的一模一样，这样你说"来源文件"
#: 的时候它才认得。两处各写一份的话，早晚改了一边忘了另一边。
AUX_HEADERS = ("来源文件", "来源行")

#: 旧名，留个别处引用不到时炸掉。新代码用 :data:`AUX_HEADERS`。
_AUX_HEADERS = AUX_HEADERS


def _cell_value(value: object) -> object:
    """把值转成 Excel 能写的类型。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def _autofit(ws, widths: list[int] | None = None) -> None:
    """按内容估算列宽。中文按 2 个字符宽度算。"""
    for idx, column_cells in enumerate(ws.iter_cols(), start=1):
        if widths is not None and idx <= len(widths) and widths[idx - 1] > 0:
            ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]
            continue
        longest = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            text = str(cell.value)
            # 中文按两个字符宽度估算
            width = sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
            longest = max(longest, width)
        ws.column_dimensions[get_column_letter(idx)].width = max(
            _MIN_COL_WIDTH, min(longest + 3, _MAX_COL_WIDTH)
        )


def write_summary(
    records: list[Record],
    fields: list[str],
    out_path: Path,
    uncertain_fields: set[str] | None = None,
) -> Path:
    """写汇总表。

    :param uncertain_fields: AI 拿不准的字段——这些列的**表头涂黄**，提醒你复核
    """
    uncertain = uncertain_fields or set()

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总表"

    headers = list(fields) + list(_AUX_HEADERS)
    ws.append(headers)
    for idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        if name in uncertain:
            # 拿不准的列，表头改成淡黄底 + 深色字，和蓝底区分开
            cell.fill = _UNCERTAIN_FILL
            cell.font = Font(name="微软雅黑", bold=True, size=11, color="9C6500")

    for record in records:
        row: list[object] = [_cell_value(record.values.get(name)) for name in fields]
        row.extend([record.file, record.row])
        ws.append(row)

    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if col_idx > len(fields):
                cell.font = _AUX_FONT
                cell.fill = _AUX_FILL
            else:
                cell.font = _BODY_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path


def write_issues(issues: list[Issue], out_path: Path, run_note: str = "") -> Path:
    """写问题清单。没有任何问题时也要生成，并在里面写一句"没问题"。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "问题清单"

    headers = ["级别", "问题", "建议怎么改", "文件", "工作表", "行号", "列名", "当前填的是"]
    ws.append(headers)
    for idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER

    if issues:
        for issue in issues:
            ws.append(
                [
                    issue.level_label,
                    issue.problem,
                    issue.suggestion,
                    issue.file,
                    issue.sheet,
                    issue.row if issue.row else "",
                    issue.column,
                    _cell_value(issue.value),
                ]
            )
        for row_idx in range(2, ws.max_row + 1):
            level_text = str(ws.cell(row=row_idx, column=1).value)
            fill, font = _LEVEL_STYLE.get(
                {"必须处理": "error", "请确认": "warning", "仅供参考": "info"}.get(level_text, "info"),
                _LEVEL_STYLE["info"],
            )
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = _BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=(col_idx in (2, 3)))
                cell.font = _BODY_FONT
            ws.cell(row=row_idx, column=1).fill = fill
            ws.cell(row=row_idx, column=1).font = font
    else:
        ws.append(["—", "没有发现问题，这份汇总可以直接用。", "", "", "", "", "", ""])
        ws.cell(row=2, column=1).font = _BODY_FONT
        ws.cell(row=2, column=2).font = Font(name="微软雅黑", size=10, bold=True, color="006100")
        ws.cell(row=2, column=2).fill = PatternFill("solid", fgColor="C6EFCE")

    ws.freeze_panes = "A2"
    _autofit(ws, widths=[10, 46, 40, 24, 14, 8, 16, 20])

    if run_note:
        note_row = ws.max_row + 2
        ws.cell(row=note_row, column=1, value=run_note).font = _AUX_FONT

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path


def default_output_paths(output_dir: Path, stamp: str | None = None) -> tuple[Path, Path]:
    """生成带日期的输出文件名，避免覆盖上一次的结果。

    ``汇总表_20260914_1530.xlsx`` / ``问题清单_20260914_1530.xlsx``
    """
    stamp = stamp or datetime.now().strftime("%Y%m%d_%H%M")
    return (
        output_dir / f"汇总表_{stamp}.xlsx",
        output_dir / f"问题清单_{stamp}.xlsx",
    )


def latest_outputs(output_dir: Path) -> tuple[Path | None, Path | None]:
    """找出最近一次生成的两个文件，给界面显示用。"""
    if not output_dir.exists():
        return None, None

    def newest(pattern: str) -> Path | None:
        items = sorted(output_dir.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
        return items[0] if items else None

    return newest("汇总表_*.xlsx"), newest("问题清单_*.xlsx")
